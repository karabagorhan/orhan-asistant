#!/usr/bin/env python3
"""Beş kalemlik, formüllü pompa bakım maliyet tablosu oluşturur."""

from argparse import ArgumentParser
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


ITEMS = [
    ("Mekanik salmastra değişimi", "adet"),
    ("Rulman değişimi", "adet"),
    ("Conta / O-ring seti", "takım"),
    ("Çark temizliği ve balans kontrolü", "adet"),
    ("Sökme-takma ve test işçiliği", "saat"),
]


def create_workbook(output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Pompa Bakım Maliyeti"
    ws.sheet_view.showGridLines = False

    navy = "17365D"
    blue = "D9EAF7"
    pale_yellow = "FFF2CC"
    gray = "666666"
    thin_gray = Side(style="thin", color="B7C9D6")

    ws.merge_cells("A1:F1")
    ws["A1"] = "POMPA BAKIM MALİYET TABLOSU — ÖRNEK"
    ws["A1"].font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=navy)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:F2")
    ws["A2"] = "Miktar ve birim fiyatları girin. Fiyatları kullanmadan önce DOĞRULAYIN. Tutarlar formülle hesaplanır; KDV hariçtir."
    ws["A2"].font = Font(name="Calibri", size=10, italic=True, color=gray)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 32

    headers = ["No", "Bakım kalemi", "Birim", "Miktar", "Birim fiyat (₺)", "Tutar (₺)"]
    header_row = 4
    for col, value in enumerate(headers, start=1):
        cell = ws.cell(header_row, col, value)
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin_gray)
    ws.row_dimensions[header_row].height = 28

    first_item_row = header_row + 1
    last_item_row = first_item_row + len(ITEMS) - 1
    for index, (description, unit) in enumerate(ITEMS, start=1):
        row = first_item_row + index - 1
        ws.cell(row, 1, index)
        ws.cell(row, 2, description)
        ws.cell(row, 3, unit)
        # Input cells stay empty: quantities and prices are not assumed.
        ws.cell(row, 4).fill = PatternFill("solid", fgColor=pale_yellow)
        ws.cell(row, 5).fill = PatternFill("solid", fgColor=pale_yellow)
        ws.cell(row, 6, f'=IF(OR(D{row}="",E{row}=""),"",D{row}*E{row})')
        for col in range(1, 7):
            cell = ws.cell(row, col)
            cell.border = Border(bottom=thin_gray)
            cell.alignment = Alignment(vertical="center", wrap_text=(col == 2))
        ws.cell(row, 1).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row, 3).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row, 4).number_format = "#,##0.00"
        ws.cell(row, 5).number_format = '₺ #,##0.00;[Red]-₺ #,##0.00'
        ws.cell(row, 6).number_format = '₺ #,##0.00;[Red]-₺ #,##0.00'

    # Reject negative quantities and prices while allowing empty input cells.
    positive_number = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
    positive_number.error = "Sıfır veya pozitif bir değer girin."
    positive_number.errorTitle = "Geçersiz değer"
    ws.add_data_validation(positive_number)
    positive_number.add(f"D{first_item_row}:E{last_item_row}")

    total_row = last_item_row + 2
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=5)
    ws.cell(total_row, 1, "GENEL TOPLAM (KDV hariç)")
    ws.cell(total_row, 6, f'=IF(COUNT(D{first_item_row}:D{last_item_row})=0,"",IF(AND(COUNT(D{first_item_row}:D{last_item_row})={len(ITEMS)},COUNT(E{first_item_row}:E{last_item_row})={len(ITEMS)}),SUM(F{first_item_row}:F{last_item_row}),"Eksik bilgi"))')
    for col in range(1, 7):
        cell = ws.cell(total_row, col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(horizontal="right" if col == 6 else "left", vertical="center")
    ws.cell(total_row, 6).number_format = '₺ #,##0.00;[Red]-₺ #,##0.00'
    ws.row_dimensions[total_row].height = 25

    ws.column_dimensions["A"].width = 7
    ws.column_dimensions["B"].width = 39
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 20
    ws.column_dimensions["F"].width = 18
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A{header_row}:F{last_item_row}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.orientation = "landscape"
    ws.print_area = f"A1:F{total_row}"

    output = output.resolve()
    wb.save(output)
    print(output)


def main() -> None:
    parser = ArgumentParser(description="Formüllü örnek pompa bakım maliyet tablosu üretir.")
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("/opt/asistan/ciktilar") / f"{date.today().isoformat()}_pompa_bakim_maliyet.xlsx",
        help="Oluşturulacak Excel dosyasının yolu",
    )
    args = parser.parse_args()
    create_workbook(args.output)


if __name__ == "__main__":
    main()
