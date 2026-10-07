"""
Hakediş Kontrol Sistemi — Götürü Bedel
Kullanım: python hakedis_kontrol.py
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import date
import argparse, os

# ── Sabitler ──────────────────────────────────────────────────────────────────
KDV_ORANI   = 0.20
STOPAJ_ORAN = 0.03   # %3 hizmet stopajı (değiştirmek için argüman kullan)

RENK = {
    "baslik"   : "1F3864",   # koyu lacivert
    "alt_baslik": "2E75B6",  # orta mavi
    "vurgu"    : "BDD7EE",   # açık mavi
    "uyari"    : "FFD966",   # sarı
    "tamam"    : "E2EFDA",   # açık yeşil
    "hata"     : "FFE0E0",   # açık kırmızı
    "yazi_beyaz": "FFFFFF",
}

def kenar(ince=True):
    s = Side(style="thin" if ince else "medium")
    return Border(left=s, right=s, top=s, bottom=s)

def stil_baslik(hucre, renk_hex, yazi_renk="FFFFFF", boyut=11, kalin=True):
    hucre.font = Font(name="Calibri", bold=kalin, size=boyut, color=yazi_renk)
    hucre.fill = PatternFill("solid", fgColor=renk_hex)
    hucre.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    hucre.border = kenar()

def stil_veri(hucre, renk_hex=None, sayi=False, yuzdeli=False):
    hucre.font = Font(name="Calibri", size=10)
    if renk_hex:
        hucre.fill = PatternFill("solid", fgColor=renk_hex)
    hucre.border = kenar()
    hucre.alignment = Alignment(vertical="center")
    if sayi:
        hucre.number_format = '#,##0.00 ₺'
    if yuzdeli:
        hucre.number_format = '0.00%'

# ── Ana Fonksiyon ──────────────────────────────────────────────────────────────
def olustur(taşeron, sozlesme_no, sozlesme_bedel, tamamlanma_yuzde,
            kdv=KDV_ORANI, stopaj=STOPAJ_ORAN, cikti_klas="/opt/asistan/ciktilar"):

    wb = openpyxl.Workbook()

    # ── Sayfa 1: Hakediş Özet ──────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Hakediş Özet"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 22

    # Başlık bloğu
    ws.merge_cells("A1:D1")
    ws["A1"] = "POLİSAN KİMYA — HAKEDİŞ KONTROL FORMU"
    stil_baslik(ws["A1"], RENK["baslik"], boyut=13)
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:D2")
    ws["A2"] = f"Tarih: {date.today().strftime('%d.%m.%Y')}  |  Taşeron: {taşeron}  |  Sözleşme No: {sozlesme_no}"
    ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="444444")
    ws["A2"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[2].height = 18

    ws.row_dimensions[3].height = 8  # boşluk

    # Alt başlık
    for col, metin in enumerate(["Kalem", "Sözleşme Değeri (₺)", "Bu Hakediş (₺)", "Kümülatif (₺)"], 1):
        h = ws.cell(row=4, column=col, value=metin)
        stil_baslik(h, RENK["alt_baslik"], boyut=10)
    ws.row_dimensions[4].height = 22

    # Veri satırları — götürü bedel yapısı
    kalemler = [
        ("Sözleşme Bedeli (KDV Hariç)", sozlesme_bedel, None, None),
    ]

    satirlar = {
        "sozlesme"   : 5,
        "tamamlanma" : 6,
        "bu_hakedis" : 7,
        "onceki"     : 8,
        "kdv_matrah" : 9,
        "kdv"        : 10,
        "stopaj"     : 11,
        "net_odeme"  : 12,
    }

    # Sözleşme bedeli
    satirlar_veri = [
        (5,  "Sözleşme Bedeli (KDV Hariç)",   sozlesme_bedel,              True),
        (6,  "Tamamlanma Yüzdesi",              tamamlanma_yuzde,            False),   # yüzde
        (7,  "Kümülatif Hakediş (Bu dahil)",    None,                        True),    # formül
        (8,  "Önceki Hakediş Toplamı",          0.0,                         True),    # kullanıcı girer
        (9,  "Bu Hakediş (KDV Matrahı)",        None,                        True),    # formül
        (10, f"KDV (%{int(kdv*100)})",          None,                        True),    # formül
        (11, f"Stopaj (-%{int(stopaj*100)})",   None,                        True),    # formül
        (12, "NET ÖDEME",                        None,                        True),    # formül
    ]

    for satir, etiket, deger, sayi_mi in satirlar_veri:
        hA = ws.cell(row=satir, column=1, value=etiket)
        hB = ws.cell(row=satir, column=2)
        hC = ws.cell(row=satir, column=3)
        hD = ws.cell(row=satir, column=4)

        r = RENK["vurgu"] if satir % 2 == 1 else None

        # A sütunu etiket
        hA.font = Font(name="Calibri", size=10, bold=(satir == 12))
        hA.border = kenar()
        hA.alignment = Alignment(vertical="center", indent=1)
        if satir == 12:
            hA.fill = PatternFill("solid", fgColor=RENK["baslik"])
            hA.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")

        ws.row_dimensions[satir].height = 20

        if satir == 5:   # Sözleşme bedeli
            hB.value = deger
            stil_veri(hB, r, sayi=True)
            ws.merge_cells(f"C{satir}:D{satir}")
            ws.cell(row=satir, column=3).value = "DOĞRULANMALI — Sözleşme ekinden kontrol edin"
            ws.cell(row=satir, column=3).fill = PatternFill("solid", fgColor=RENK["uyari"])
            ws.cell(row=satir, column=3).font = Font(name="Calibri", size=9, italic=True, color="7F6000")
            ws.cell(row=satir, column=3).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row=satir, column=3).border = kenar()

        elif satir == 6:  # Tamamlanma yüzdesi
            hB.value = deger
            hB.number_format = '0.00%'
            hB.font = Font(name="Calibri", size=10, bold=True, color="1F3864")
            hB.fill = PatternFill("solid", fgColor=RENK["tamam"])
            hB.border = kenar()
            hB.alignment = Alignment(horizontal="center", vertical="center")
            ws.merge_cells(f"C{satir}:D{satir}")
            ws.cell(row=satir, column=3).value = "Saha kontrol yüzdesi — saha ekibinden teyit alın"
            ws.cell(row=satir, column=3).fill = PatternFill("solid", fgColor=RENK["uyari"])
            ws.cell(row=satir, column=3).font = Font(name="Calibri", size=9, italic=True, color="7F6000")
            ws.cell(row=satir, column=3).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row=satir, column=3).border = kenar()

        elif satir == 7:  # Kümülatif = sözleşme × tamamlanma
            hB.value = f"=B5*B6"
            stil_veri(hB, RENK["tamam"], sayi=True)
            hB.font = Font(name="Calibri", size=10, bold=True)
            for c in [hC, hD]:
                stil_veri(c, r)

        elif satir == 8:  # Önceki ödemeler — kullanıcı girer
            hB.value = 0.0
            hB.fill = PatternFill("solid", fgColor="FFF2CC")
            hB.font = Font(name="Calibri", size=10)
            hB.border = kenar()
            hB.number_format = '#,##0.00 ₺'
            hB.alignment = Alignment(vertical="center")
            ws.merge_cells(f"C{satir}:D{satir}")
            ws.cell(row=satir, column=3).value = "▶ Bu hücreye önceki toplam ödemeyi girin"
            ws.cell(row=satir, column=3).fill = PatternFill("solid", fgColor="FFF2CC")
            ws.cell(row=satir, column=3).font = Font(name="Calibri", size=9, color="7F6000")
            ws.cell(row=satir, column=3).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row=satir, column=3).border = kenar()

        elif satir == 9:  # Bu hakediş matrahı
            hB.value = "=B7-B8"
            stil_veri(hB, r, sayi=True)
            for c in [hC, hD]:
                stil_veri(c, r)

        elif satir == 10:  # KDV
            hB.value = f"=B9*{kdv}"
            stil_veri(hB, r, sayi=True)
            for c in [hC, hD]:
                stil_veri(c, r)

        elif satir == 11:  # Stopaj (negatif)
            hB.value = f"=-B9*{stopaj}"
            hB.fill = PatternFill("solid", fgColor=RENK["hata"])
            hB.font = Font(name="Calibri", size=10, color="9C0006")
            hB.border = kenar()
            hB.number_format = '#,##0.00 ₺'
            hB.alignment = Alignment(vertical="center")
            for c in [hC, hD]:
                stil_veri(c, r)

        elif satir == 12:  # Net ödeme
            hB.value = "=B9+B10+B11"
            hB.fill = PatternFill("solid", fgColor=RENK["baslik"])
            hB.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            hB.border = kenar(ince=False)
            hB.number_format = '#,##0.00 ₺'
            hB.alignment = Alignment(horizontal="center", vertical="center")
            for c in [hC, hD]:
                c.fill = PatternFill("solid", fgColor=RENK["baslik"])
                c.border = kenar()

    # ── Sayfa 2: Malzeme Kalemi ────────────────────────────────────────────────
    ws2 = wb.create_sheet("Malzeme Kalemleri")
    ws2.sheet_view.showGridLines = False

    sutunlar = [("A", 35, "Malzeme / Ekipman Adı"),
                ("B", 12, "Birim"),
                ("C", 12, "Miktar"),
                ("D", 18, "Birim Fiyat (₺)"),
                ("E", 18, "Tutar (₺)"),
                ("F", 10, "KDV %"),
                ("G", 18, "KDV Tutarı (₺)"),
                ("H", 20, "Not / Kaynak")]

    for harf, genislik, baslik in sutunlar:
        ws2.column_dimensions[harf].width = genislik
        h = ws2[f"{harf}1"]
        h.value = baslik
        stil_baslik(h, RENK["alt_baslik"], boyut=10)

    ws2.row_dimensions[1].height = 22

    # 20 boş satır + formüller
    for i in range(2, 22):
        renk = RENK["vurgu"] if i % 2 == 0 else None
        for harf in ["A", "B", "C", "D", "E", "F", "G", "H"]:
            c = ws2[f"{harf}{i}"]
            stil_veri(c, renk, sayi=(harf in ["D", "E", "G"]))
            ws2.row_dimensions[i].height = 18
        # E = miktar × birim fiyat
        ws2[f"E{i}"].value = f"=IF(C{i}*D{i}=0,\"\",C{i}*D{i})"
        ws2[f"E{i}"].number_format = '#,##0.00 ₺'
        # F varsayılan KDV
        ws2[f"F{i}"].value = int(kdv * 100)
        ws2[f"F{i}"].number_format = '0"%"'
        # G = KDV
        ws2[f"G{i}"].value = f'=IF(E{i}="","",E{i}*F{i}/100)'
        ws2[f"G{i}"].number_format = '#,##0.00 ₺'

    # Toplam satırı
    toplam_r = 23
    ws2.merge_cells(f"A{toplam_r}:D{toplam_r}")
    t = ws2[f"A{toplam_r}"]
    t.value = "TOPLAM"
    stil_baslik(t, RENK["baslik"], boyut=11)
    ws2.row_dimensions[toplam_r].height = 22

    for harf, formul in [("E", f"=SUM(E2:E22)"), ("G", f"=SUM(G2:G22)")]:
        c = ws2[f"{harf}{toplam_r}"]
        c.value = formul
        c.number_format = '#,##0.00 ₺'
        c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=RENK["baslik"])
        c.border = kenar(ince=False)
        c.alignment = Alignment(horizontal="center", vertical="center")

    # ── Kaydet ────────────────────────────────────────────────────────────────
    dosya_adi = f"{date.today().strftime('%Y-%m-%d')}_hakedis_{taşeron.replace(' ','_')}.xlsx"
    tam_yol = os.path.join(cikti_klas, dosya_adi)
    os.makedirs(cikti_klas, exist_ok=True)
    wb.save(tam_yol)
    return tam_yol


# ── CLI ────────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Hakediş Kontrol Excel Üretici")
    ap.add_argument("--taseron",   default="Taşeron Adı")
    ap.add_argument("--sozlesme",  default="SZL-2025-001")
    ap.add_argument("--bedel",     type=float, default=1_000_000.0)
    ap.add_argument("--yuzde",     type=float, default=0.50,
                    help="Tamamlanma yüzdesi (0-1 arası, örn: 0.65)")
    ap.add_argument("--kdv",       type=float, default=0.20)
    ap.add_argument("--stopaj",    type=float, default=0.03)
    ap.add_argument("--cikti",     default="/opt/asistan/ciktilar")
    args = ap.parse_args()

    yol = olustur(
        taşeron          = args.taseron,
        sozlesme_no      = args.sozlesme,
        sozlesme_bedel   = args.bedel,
        tamamlanma_yuzde = args.yuzde,
        kdv              = args.kdv,
        stopaj           = args.stopaj,
        cikti_klas       = args.cikti,
    )
    print(f"✅ Dosya oluşturuldu:\n   {yol}")
